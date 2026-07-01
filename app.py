import os
import re
import uuid
import datetime
import traceback
import json
import time
from typing import List
from dotenv import load_dotenv
from os import getenv

from flask import (
    Flask,
    request,
    jsonify,
    session,
    render_template,
    send_from_directory,
)
from flask_cors import CORS
from flask_session import Session
from werkzeug.utils import secure_filename

from langchain_google_genai import ChatGoogleGenerativeAI
from langgraph.prebuilt import create_react_agent
from langchain_core.messages import HumanMessage, AIMessage
from langchain_core.tools import BaseTool
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder

from langchain_community.utilities.sql_database import SQLDatabase
from langchain_community.tools.sql_database.tool import QuerySQLDataBaseTool
from sqlalchemy import create_engine, text

from PyPDF2 import PdfReader
from docx import Document
import csv
import openpyxl
import xlrd

# --- CONFIGURATION ---
load_dotenv()

app = Flask(__name__, static_folder="static", template_folder="templates")
CORS(app, supports_credentials=True)
app.config['SECRET_KEY'] = os.getenv("SECRET_KEY", "dev_secret")
app.config['SESSION_TYPE'] = 'filesystem'
app.config['SESSION_PERMANENT'] = False
Session(app)

UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
ALLOWED_EXTENSIONS = set(
    os.getenv("ALLOWED_EXTENSIONS", "pdf,docx,txt,csv,xls,xlsx,png,jpg,jpeg").split(",")
)

SENSITIVE_COLUMNS = [
    col.strip().lower()
    for col in os.getenv("SENSITIVE_COLUMNS", "").split(",")
    if col.strip()
]
BLOCKED_QUERIES = [
    word.strip().lower()
    for word in os.getenv("BLOCKED_QUERIES", "").split(",")
    if word.strip()
]

app.config['SQLALCHEMY_ENGINE'] = None

_global_store = {
    "messages": [],
}

sql_agent_executor = None
fallback_llm = None
sql_llm = None
db_schema_global = ""

IST = datetime.timezone(datetime.timedelta(hours=5, minutes=30))

# --- ENV HELPER ---
def get_env(key, default=None):
    """Get environment variable with multiple fallback strategies."""
    return os.environ.get(key) or os.getenv(key) or default

# --- UTILITIES ---
def allowed_file(filename: str) -> bool:
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def get_tables_from_sql_file(file_path: str) -> List[str]:
    table_names = []
    table_regex = re.compile(r"CREATE TABLE `?(\w+)`?", re.IGNORECASE)
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
            matches = table_regex.finditer(content)
            for match in matches:
                table_names.append(match.group(1))
    except Exception:
        return []
    return sorted(list(set(table_names)))

def extract_text_from_file(filepath: str) -> str:
    ext = filepath.split('.')[-1].lower()
    text = ""
    try:
        if ext == "pdf":
            with open(filepath, "rb") as f:
                reader = PdfReader(f)
                text = "\n".join([(page.extract_text() or "") for page in reader.pages])
        elif ext in ["doc", "docx"]:
            doc = Document(filepath)
            text = "\n".join([p.text for p in doc.paragraphs])
        elif ext == "txt":
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext == "csv":
            with open(filepath, newline='', encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                text = "\n".join([",".join([str(c) for c in row]) for row in reader])
        elif ext in ["xls", "xlsx"]:
            if ext == "xlsx":
                wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
                sheet = wb.active
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(",".join([str(cell) if cell is not None else "" for cell in row]))
                text = "\n".join(rows)
            else:
                wb = xlrd.open_workbook(filepath)
                sheet = wb.sheet_by_index(0)
                rows = []
                for i in range(sheet.nrows):
                    rows.append(",".join([str(cell) if cell is not None else "" for cell in sheet.row_values(i)]))
                text = "\n".join(rows)
        else:
            text = "Unsupported file format."
    except Exception as e:
        text = f"Error reading file: {e}"
    return text.strip()

def detect_sensitive_query(text: str) -> bool:
    query_lower = (text or "").lower()
    if any(word in query_lower for word in BLOCKED_QUERIES):
        return True
    if any(re.search(rf"\b{re.escape(col)}\b", query_lower) for col in SENSITIVE_COLUMNS):
        if "count" in query_lower or "how many" in query_lower:
            return False
        return True
    return False

def get_fallback_response(prompt: str) -> str:
    global fallback_llm
    if not fallback_llm:
        return "Error: LLM not initialized."
    try:
        system = "You are a helpful assistant. If the user is asking about SQL data, answer based only on the conversation so far."
        resp = fallback_llm.invoke([HumanMessage(content=system + "\n\n" + prompt)])
        if isinstance(resp, AIMessage):
            return resp.content if isinstance(resp.content, str) else str(resp.content)
        return str(resp)
    except Exception as e:
        print("Fallback LLM error:", e, traceback.format_exc())
        return f"Error in fallback LLM: {e}"

# --- INITIALIZATION ---
def initialize_agents():
    global sql_agent_executor, fallback_llm, sql_llm, db_schema_global

    GEMINI_API_KEY = get_env("GEMINI_API_KEY") or get_env("GOOGLE_API_KEY") or ""
    GEMINI_MODEL = get_env("GEMINI_MODEL") or "gemini-2.5-pro"

    try:
        sql_llm = ChatGoogleGenerativeAI(model=GEMINI_MODEL, google_api_key=GEMINI_API_KEY, temperature=0)
        fallback_llm = ChatGoogleGenerativeAI(model=GEMINI_MODEL, google_api_key=GEMINI_API_KEY, temperature=0.2)
        print("LLMs initialized successfully")
    except Exception as e:
        print(" LLM initialization error:", e)
        sql_llm = None
        fallback_llm = None

    db_uri = get_env("DATABASE_URL")
    if not db_uri:
        print("--- WARNING: DATABASE_URL not found. SQL Agent disabled (will use fallback responses). ---")
        app.config['SQLALCHEMY_ENGINE'] = None
        sql_agent_executor = None
        return

    db = None
    db_schema = ""
    try:
        # Try multiple possible SQL schema filenames
        sql_schema_file = None
        possible_files = ["TICKETING.sql", "TICKETING_BKP.sql", "ticketing.sql", "schema.sql"]
        for f in possible_files:
            if os.path.exists(f):
                sql_schema_file = f
                break

        if sql_schema_file:
            print(f" Found SQL schema file: {sql_schema_file}")
            all_tables = get_tables_from_sql_file(sql_schema_file)
        else:
            print(" No SQL schema file found. Will include all tables.")
            all_tables = []

        exclude = {'tbl_chat_messages', 'tbl_chats'}
        include_only_these_tables = [t for t in all_tables if t not in exclude]

        if include_only_these_tables:
            db = SQLDatabase.from_uri(db_uri, sample_rows_in_table_info=0, include_tables=include_only_these_tables)
        else:
            db = SQLDatabase.from_uri(db_uri, sample_rows_in_table_info=0)

        db_schema = db.get_table_info()
        db_schema_global = str(db_schema)
        print(f" Database schema loaded. Tables: {include_only_these_tables or 'all'}")
    except Exception as e:
        print(" Error building SQLDatabase from uri:", e, traceback.format_exc())
        db = None
        db_schema = ""

    safe_schema = str(db_schema).replace("{", "{{").replace("}", "}}")
    system_prompt_string = f"""
You are a highly methodical and precise MySQL data analyst for a banking ticket management system.

Your job:
- Understand the user's question.
- Decide which tables and columns to query from the schema.
- Write a single, syntactically correct MySQL query.
- Use the SQL tool to execute it.
- Then explain the result in simple, clear English for a non-technical banking manager.

Very important rules:
1. Always answer based on the database schema below, not on guesswork.
2. If the question is about counts, totals, or group-by, write and run an appropriate aggregate SQL query.
3. **CRITICAL:** Do NOT add filters (like 'status', 'date', or 'active') unless the user EXPLICITLY asks for them.
4. Never return raw SQL as the final answer. Always convert results into a human-readable explanation.
5. Format final answers in **Markdown**:
   - If the answer is a single value or short explanation, respond as one or two sentences.
   - If the result is tabular, present a clear **Markdown table**.
   - **IMPORTANT:** Ensure every row is on a new line. Do NOT concatenate rows.
6. If the data truly does not exist in the schema, say so clearly.

Database schema:
{safe_schema}
""".strip()

    agent_prompt = None
    try:
        agent_prompt = ChatPromptTemplate.from_messages([
            ("system", system_prompt_string),
            MessagesPlaceholder(variable_name="messages"),
        ])
    except Exception as e:
        print(" Prompt template creation error:", e, traceback.format_exc())

    try:
        tools: List[BaseTool] = []
        if db is not None:
            tools = [QuerySQLDataBaseTool(db=db)]

        if sql_llm and agent_prompt:
            sql_agent_executor = create_react_agent(model=sql_llm, tools=tools, prompt=agent_prompt)
            print("[OK] SQL agent created successfully.")
        else:
            sql_agent_executor = None
            print(" SQL agent NOT created (missing LLM or prompt).")
    except Exception as e:
        print(" Error creating react agent:", e, traceback.format_exc())
        sql_agent_executor = None

    try:
        engine = create_engine(db_uri, pool_pre_ping=True, pool_recycle=300)
        app.config['SQLALCHEMY_ENGINE'] = engine
        print(" SQLAlchemy engine created successfully.")
    except Exception as e:
        print(" Error creating SQLAlchemy engine:", e, traceback.format_exc())
        app.config['SQLALCHEMY_ENGINE'] = None


def initialize_agents_with_retry(max_retries=5, delay=5):
    """Initialize agents with retry logic for cloud deployment."""
    global sql_agent_executor, fallback_llm, sql_llm, db_schema_global

    for attempt in range(max_retries):
        try:
            initialize_agents()
            if app.config.get('SQLALCHEMY_ENGINE') is not None:
                print(" Database connected successfully on attempt", attempt + 1)
                return True
            print(f" DB not ready, retrying in {delay}s... ({attempt + 1}/{max_retries})")
            time.sleep(delay)
        except Exception as e:
            print(f" Init attempt {attempt + 1} failed: {e}")
            if attempt < max_retries - 1:
                time.sleep(delay)

    print("[WARNING] Running without database - fallback mode only")
    return False

# --- ROUTES ---
@app.route('/')
def index():
    template_name = "index.html"
    template_path = os.path.join(app.root_path, "templates", template_name)
    file_root = os.path.join(app.root_path, template_name)
    if os.path.exists(template_path): return render_template(template_name)
    if os.path.exists(file_root): return send_from_directory(app.root_path, template_name)
    return f"Missing {template_name} in templates/ or project root.", 404

@app.route('/api/upload', methods=['POST'])
def upload_file_route():
    if 'file' not in request.files: return jsonify({"error": "No file part"}), 400
    f = request.files['file']
    if f.filename == "": return jsonify({"error": "No selected file"}), 400
    if not allowed_file(f.filename): return jsonify({"error": "Invalid file type"}), 400

    filename = secure_filename(f.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    try:
        f.save(filepath)
        extracted_text = extract_text_from_file(filepath)
        session['uploaded_context'] = extracted_text
        session.modified = True
        return jsonify({"message": "uploaded", "filename": filename, "content": extracted_text})
    except Exception as e:
        print("Upload error:", e, traceback.format_exc())
        return jsonify({"error": "Upload failed"}), 500

@app.route('/api/history', methods=['GET'])
def get_history_route():
    page = int(request.args.get("page", 1))
    per_page = 50
    offset = (page - 1) * per_page
    engine = app.config.get('SQLALCHEMY_ENGINE')

    if engine:
        try:
            with engine.connect() as conn:
                rows = conn.execute(
                    text("""
                        SELECT message_id AS id, role, message, timestamp
                        FROM tbl_chat_messages
                        ORDER BY timestamp DESC
                        LIMIT :limit OFFSET :offset
                    """), {"limit": per_page, "offset": offset}
                ).all()

                messages = []
                for r in rows:
                    ts = r.timestamp
                    ts_iso = ts.isoformat() if isinstance(ts, datetime.datetime) else str(ts)
                    messages.append({
                        "id": str(r.id),
                        "role": r.role,
                        "message": r.message,
                        "timestamp": ts_iso
                    })
                return jsonify(messages)
        except Exception as e:
            print("DB /api/history error:", e, traceback.format_exc())

    msgs = _global_store["messages"]
    start = offset
    end = offset + per_page
    sliced = msgs[::-1][start:end] 
    return jsonify(sliced)

@app.route('/api/history/date', methods=['GET'])
def get_history_by_date_route():
    date_str = request.args.get("date")
    if not date_str: return jsonify({"reset": True})

    engine = app.config.get('SQLALCHEMY_ENGINE')
    if engine:
        try:
            with engine.connect() as conn:
                start_of_day = f"{date_str} 00:00:00"
                end_of_day = f"{date_str} 23:59:59"

                target_row = conn.execute(
                    text("""
                        SELECT message_id, timestamp
                        FROM tbl_chat_messages 
                        WHERE timestamp >= :start AND timestamp <= :end
                        ORDER BY timestamp ASC
                        LIMIT 1
                    """), {"start": start_of_day, "end": end_of_day}
                ).first()

                if not target_row: return jsonify({"error": f"No messages found for {date_str}"})

                target_msg_id = str(target_row.message_id)
                target_ts = target_row.timestamp

                count_row = conn.execute(
                    text("SELECT COUNT(*) as cnt FROM tbl_chat_messages WHERE timestamp > :ts"), 
                    {"ts": target_ts}
                ).first()

                newer_count = count_row.cnt if count_row else 0
                per_page = 50
                target_page = (newer_count // per_page) + 1
                offset = (target_page - 1) * per_page

                rows = conn.execute(
                    text("""
                        SELECT message_id AS id, role, message, timestamp
                        FROM tbl_chat_messages
                        ORDER BY timestamp DESC
                        LIMIT :limit OFFSET :offset
                    """), {"limit": per_page, "offset": offset}
                ).all()

                messages = []
                for r in rows:
                    ts = r.timestamp
                    ts_iso = ts.isoformat() if isinstance(ts, datetime.datetime) else str(ts)
                    messages.append({
                        "id": str(r.id),
                        "role": r.role,
                        "message": r.message,
                        "timestamp": ts_iso
                    })

                return jsonify({
                    "messages": messages, 
                    "new_current_page": target_page,
                    "target_message_id": target_msg_id
                })

        except Exception as e:
            print("DB /api/history/date error:", e, traceback.format_exc())
            return jsonify({"error": "Database error during date jump"})

    return jsonify({"error": "DB connection unavailable"})

@app.route('/api/message/delete', methods=['POST'])
def delete_message_route():
    data = request.get_json() or {}
    message_id = data.get("message_id")
    if not message_id: return jsonify({"error": "message_id required"}), 400

    engine = app.config.get('SQLALCHEMY_ENGINE')
    deleted_ids = []

    if engine:
        try:
            with engine.connect() as conn:
                row = conn.execute(
                    text("SELECT message_id, role, timestamp FROM tbl_chat_messages WHERE message_id = :mid"),
                    {"mid": message_id},
                ).first()
                if not row: return jsonify({"error": "not found"}), 404

                deleted_ids.append(str(row.message_id))

                if row.role == 'user':
                    next_row = conn.execute(
                        text("""
                            SELECT message_id FROM tbl_chat_messages
                            WHERE timestamp > :ts
                            ORDER BY timestamp ASC
                            LIMIT 1
                        """), {"ts": row.timestamp},
                    ).first()
                    if next_row:
                        deleted_ids.append(str(next_row.message_id))
                        conn.execute(text("DELETE FROM tbl_chat_messages WHERE message_id = :mid"), {"mid": next_row.message_id})

                conn.execute(text("DELETE FROM tbl_chat_messages WHERE message_id = :mid"), {"mid": row.message_id})
                conn.commit()
                return jsonify({"deleted_ids": deleted_ids})
        except Exception as e:
            print("/api/message/delete DB error:", e, traceback.format_exc())
            return jsonify({"error": "failed"}), 500

    remaining = []
    for m in _global_store["messages"]:
        if m["id"] == message_id:
            deleted_ids.append(message_id)
            continue
        remaining.append(m)
    _global_store["messages"] = remaining
    return jsonify({"deleted_ids": deleted_ids})

@app.route('/api/chat', methods=['POST'])
def query_route():
    global sql_agent_executor
    data = request.get_json() or {}
    message = (data.get("message") or "").strip()
    file_content = data.get("file_content")

    if not message and not file_content:
        return jsonify({"response": "Please enter a valid query."}), 400

    engine = app.config.get('SQLALCHEMY_ENGINE')
    now_ist = datetime.datetime.now(IST)

    if file_content:
        session['uploaded_context'] = file_content
        session.modified = True

    user_message_id = str(uuid.uuid4())
    try:
        if engine:
            with engine.connect() as conn:
                conn.execute(
                    text("""
                        INSERT INTO tbl_chat_messages (message_id, role, message, timestamp)
                        VALUES (:mid, 'user', :msg, :ts)
                    """), {"mid": user_message_id, "msg": message, "ts": now_ist}
                )
                conn.commit()
        else:
            _global_store["messages"].append({
                "id": user_message_id, "role": "user", "message": message, "timestamp": now_ist.isoformat()
            })
    except Exception as e:
        print("Error persisting user message:", e, traceback.format_exc())

    if detect_sensitive_query(message):
        resp_text = "Sorry — I can't provide sensitive personal data. You can ask for aggregated counts though."
        bot_message_id = str(uuid.uuid4())
        try:
            if engine:
                with engine.connect() as conn:
                    conn.execute(
                        text("""
                            INSERT INTO tbl_chat_messages (message_id, role, message, timestamp)
                            VALUES (:mid, 'bot', :msg, :ts)
                        """), {"mid": bot_message_id, "msg": resp_text, "ts": now_ist}
                    )
                    conn.commit()
            else:
                _global_store["messages"].append({
                    "id": bot_message_id, "role": "bot", "message": resp_text, "timestamp": now_ist.isoformat()
                })
        except Exception as e:
            print("Error persisting sensitive reply:", e, traceback.format_exc())
        return jsonify({"response": resp_text, "user_message_id": user_message_id, "bot_message_id": bot_message_id})

    conversation_history: List[HumanMessage | AIMessage] = []
    try:
        if engine:
            with engine.connect() as conn:
                rows = conn.execute(text("SELECT role, message, timestamp FROM tbl_chat_messages ORDER BY timestamp")).all()
                for r in rows:
                    if r.role == 'user': conversation_history.append(HumanMessage(content=r.message))
                    else: conversation_history.append(AIMessage(content=r.message))
        else:
            for m in sorted(_global_store["messages"], key=lambda x: x["timestamp"]):
                if m["role"] == "user": conversation_history.append(HumanMessage(content=m["message"]))
                else: conversation_history.append(AIMessage(content=m["message"]))
    except Exception as e:
        print("Error building conversation_history:", e, traceback.format_exc())
        conversation_history = [HumanMessage(content=message)]

    ai_response = "An error occurred."
    agent_error = None

    if sql_agent_executor is not None:
        try:
            response_graph = sql_agent_executor.invoke({"messages": conversation_history})
            try:
                if isinstance(response_graph, dict) and "messages" in response_graph:
                    final_msg = response_graph["messages"][-1]
                    raw = final_msg.content if isinstance(final_msg, AIMessage) else final_msg
                elif isinstance(response_graph, AIMessage):
                    raw = response_graph.content
                else:
                    raw = response_graph

                if isinstance(raw, str): ai_response = raw
                elif isinstance(raw, (list, dict)): ai_response = json.dumps(raw, ensure_ascii=False)
                else: ai_response = str(raw)
            except Exception as e:
                print("Error parsing agent response:", e, traceback.format_exc())
                ai_response = str(response_graph)
        except Exception as e:
            print("Agent invocation error:", e, traceback.format_exc())
            agent_error = e
    else:
        agent_error = RuntimeError("SQL agent not available")

    if agent_error is not None:
        if fallback_llm: ai_response = get_fallback_response(message)
        else: ai_response = f"Agent error: {agent_error}"

    try:
        if isinstance(ai_response, (list, dict)):
            if isinstance(ai_response, list) and len(ai_response) > 0:
                if isinstance(ai_response[0], dict) and "text" in ai_response[0]: ai_response = ai_response[0]["text"]
            elif isinstance(ai_response, dict) and "text" in ai_response: ai_response = ai_response["text"]
            else: ai_response = json.dumps(ai_response, ensure_ascii=False)
        elif isinstance(ai_response, str):
            clean_str = ai_response.strip()
            if clean_str.startswith("```"):
                clean_str = clean_str[7:] if clean_str.startswith("```json") else clean_str[3:]
                if clean_str.endswith("```"): clean_str = clean_str[:-3]
                clean_str = clean_str.strip()
            if (clean_str.startswith("[") and clean_str.endswith("]")) or (clean_str.startswith("{") and clean_str.endswith("}")):
                try:
                    parsed = json.loads(clean_str)
                    if isinstance(parsed, list) and len(parsed) > 0:
                        if isinstance(parsed[0], dict) and "text" in parsed[0]: ai_response = parsed[0]["text"]
                    elif isinstance(parsed, dict) and "text" in parsed: ai_response = parsed["text"]
                except json.JSONDecodeError: pass
    except Exception as e:
        print("Error cleaning up JSON response:", e)

    if not isinstance(ai_response, str): ai_response = str(ai_response)
    if not ai_response or not ai_response.strip():
        print("WARNING: Agent returned an empty response. Using fallback.")
        ai_response = "I apologize, but I was unable to generate a response to that question based on the available data."

    bot_message_id = str(uuid.uuid4())
    bot_ts = datetime.datetime.now(IST)

    try:
        if engine:
            with engine.connect() as conn:
                conn.execute(
                    text("""
                        INSERT INTO tbl_chat_messages (message_id, role, message, timestamp)
                        VALUES (:mid, 'bot', :msg, :ts)
                    """), {"mid": bot_message_id, "msg": ai_response, "ts": bot_ts}
                )
                conn.commit()
        else:
            _global_store["messages"].append({
                "id": bot_message_id, "role": "bot", "message": ai_response, "timestamp": bot_ts.isoformat()
            })
    except Exception as e:
        print("Error persisting bot response:", e, traceback.format_exc())

    return jsonify({"response": ai_response, "user_message_id": user_message_id, "bot_message_id": bot_message_id})


# Health check endpoint for deployment platforms
@app.route('/health')
def health_check():
    return jsonify({
        "status": "healthy",
        "db_connected": app.config.get('SQLALCHEMY_ENGINE') is not None,
        "llm_initialized": fallback_llm is not None,
        "sql_agent_ready": sql_agent_executor is not None,
        "timestamp": datetime.datetime.now(IST).isoformat()
    })


# Initialize with retry for cloud deployment
initialize_agents_with_retry()

if __name__ == '__main__':
    port = int(os.getenv("PORT", 5000))
    print("Starting app on port", port)
    app.run(host="0.0.0.0", port=port, debug=False)